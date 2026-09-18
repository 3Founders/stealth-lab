"""DOCX adapter (python-docx). Provenance unit is paragraph/table index
within the document -- DOCX has no stable page numbers outside a
rendered layout, so paragraph identity is the honest choice (Prompt 2:
"DOCX section -> document + paragraph/section identity")."""
from __future__ import annotations

import hashlib
import re
from datetime import datetime, timezone

from app.services.ingestion_sources.canonical import (
    CanonicalDocument,
    SectionElement,
    SourceLocation,
    TableElement,
)
from app.services.ingestion_sources.document_adapter import (
    AdapterNotApplicable,
    DocumentLocator,
    DocumentSourceAdapter,
    RawDocument,
)

_HEADING_STYLE_RE = re.compile(r"^Heading\s*(\d)$", re.IGNORECASE)


def _source_id(uri: str) -> str:
    return hashlib.sha256(f"docx\0{uri}".encode("utf-8")).hexdigest()[:32]


class DocxAdapter(DocumentSourceAdapter):
    source_type = "document_docx"

    def can_handle(self, locator: DocumentLocator) -> bool:
        if locator.content_type_hint in (
            "docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ):
            return True
        name = (locator.filename or locator.local_path or locator.uri or "").lower()
        return name.endswith(".docx")

    def fetch(self, locator: DocumentLocator) -> RawDocument:
        if locator.raw_bytes is not None:
            data = locator.raw_bytes
        elif locator.local_path:
            with open(locator.local_path, "rb") as fh:
                data = fh.read()
        else:
            raise AdapterNotApplicable("DocxAdapter.fetch requires local_path or raw_bytes")
        uri = locator.uri or locator.local_path or locator.filename or "unknown.docx"
        return RawDocument(
            source_type=self.source_type, uri=uri, content=data,
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            repository=locator.repository, path=locator.path, commit=locator.commit,
            fetched_at=datetime.now(timezone.utc), metadata=dict(locator.metadata),
        )

    def normalize(self, raw: RawDocument) -> CanonicalDocument:
        import io

        import docx

        warnings: list[tuple[str, str]] = []
        sections: list[SectionElement] = []
        tables: list[TableElement] = []
        markdown_parts: list[str] = []

        try:
            document = docx.Document(io.BytesIO(raw.content))
        except Exception as exc:
            doc = CanonicalDocument(
                source_id=_source_id(raw.uri), source_type=self.source_type,
                source_uri=raw.uri, title=raw.path or raw.uri, canonical_markdown="",
                repository=raw.repository, path=raw.path, version=raw.commit,
                metadata=dict(raw.metadata),
            )
            return doc.with_warning("quarantine:docx_open_failed", f"{type(exc).__name__}: {exc}")

        title: str | None = None
        # python-docx exposes body children only via the underlying XML
        # element order; walking `document.element.body` in document
        # order is the only way to interleave paragraphs and tables
        # correctly rather than dumping all paragraphs then all tables.
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        def _iter_block_items(parent):
            for child in parent.element.body.iterchildren():
                if child.tag.endswith("}p"):
                    yield Paragraph(child, parent)
                elif child.tag.endswith("}tbl"):
                    yield Table(child, parent)

        paragraph_index = -1
        table_index = -1
        for item in _iter_block_items(document):
            if isinstance(item, Paragraph):
                paragraph_index += 1
                text = item.text.strip()
                if not text:
                    continue
                style_name = (item.style.name if item.style else "") or ""
                heading_match = _HEADING_STYLE_RE.match(style_name)
                if heading_match:
                    level = int(heading_match.group(1))
                    if title is None:
                        title = text
                    sections.append(
                        SectionElement(
                            heading=text, level=level, position=paragraph_index, text=text,
                            location=SourceLocation(paragraph_index=paragraph_index),
                        )
                    )
                    markdown_parts.append(f"{'#' * level} {text}")
                elif style_name.lower() == "title" and title is None:
                    title = text
                    markdown_parts.append(f"# {text}")
                else:
                    markdown_parts.append(text)
            elif isinstance(item, Table):
                table_index += 1
                try:
                    rows = [[cell.text.strip() for cell in row.cells] for row in item.rows]
                except Exception as exc:
                    warnings.append(("warning:table_read_failed", f"table {table_index}: {type(exc).__name__}: {exc}"))
                    continue
                if not rows:
                    continue
                headers, body = rows[0], tuple(tuple(r) for r in rows[1:])
                tables.append(
                    TableElement(
                        position=table_index, headers=tuple(headers), rows=body,
                        location=SourceLocation(table_index=table_index),
                    )
                )
                md_rows = ["| " + " | ".join(headers) + " |",
                           "|" + "|".join(" --- " for _ in headers) + "|"]
                md_rows.extend("| " + " | ".join(row) + " |" for row in body)
                markdown_parts.append("\n".join(md_rows))

        if not markdown_parts:
            warnings.append(("quarantine:no_extractable_content", "document had no non-empty paragraphs or tables"))

        canonical_markdown = "\n\n".join(markdown_parts) + ("\n" if markdown_parts else "")

        doc = CanonicalDocument(
            source_id=_source_id(raw.uri),
            source_type=self.source_type,
            source_uri=raw.uri,
            title=title or raw.path or raw.uri,
            canonical_markdown=canonical_markdown,
            sections=tuple(sections),
            tables=tuple(tables),
            repository=raw.repository,
            path=raw.path,
            version=raw.commit,
            metadata=dict(raw.metadata),
        )
        for code, message in warnings:
            doc = doc.with_warning(code, message)
        return doc
