"""
The deterministic half of the PDF -> Excel workflow.

No pdfplumber here: type inference, template matching, and mapping
application are pure functions over already-extracted cells, which is what
makes stages 4 and 6 exact and the six-stage chain viable at all.
"""
from __future__ import annotations

import pytest

from app.runners.base import RunContext, RunnerError
from app.runners.tables import (
    apply_column_mapping,
    map_to_schema_template,
    validate_types,
    write_xlsx,
)
from app.services.criteria import evaluate_criteria


@pytest.fixture
def ctx(tmp_path) -> RunContext:
    return RunContext(workdir=tmp_path, node_ref="stage")


GRID = [
    ["A-100", "Widget", "3", "$12.50", "2026-01-04"],
    ["A-101", "Gadget", "10", "$1,204.00", "2026-02-11"],
]
HEADER = ["SKU", "Description", "Qty", "Unit Price", "Ordered"]


def test_column_types_are_inferred_from_every_value(ctx):
    out = validate_types({"header": HEADER, "grid": GRID}, ctx)
    assert [c["type"] for c in out["columns"]] == [
        "string", "string", "integer", "number", "date"
    ]
    assert out["errors"] == []


def test_currency_and_thousands_separators_are_coerced(ctx):
    out = validate_types({"header": HEADER, "grid": GRID}, ctx)
    assert out["typed_grid"][1][3] == 1204.0
    assert out["typed_grid"][0][2] == 3


def test_dates_are_normalised_to_iso(ctx):
    out = validate_types({"header": HEADER, "grid": GRID}, ctx)
    assert out["typed_grid"][0][4] == "2026-01-04"


def test_a_mixed_column_widens_to_string_rather_than_failing(ctx):
    grid = [["1"], ["not a number"]]
    out = validate_types({"header": ["Value"], "grid": grid}, ctx)

    # Inferring per column rather than per cell makes coercion total: there
    # is no "this cell wouldn't convert" failure mode to report.
    assert out["columns"][0]["type"] == "string"
    assert out["errors"] == []


def test_a_ragged_row_is_an_error_and_is_excluded(ctx):
    grid = GRID + [["A-102", "Short row"]]
    out = validate_types({"header": HEADER, "grid": grid}, ctx)

    assert len(out["errors"]) == 1
    assert "2 cells" in out["errors"][0]
    assert len(out["typed_grid"]) == 2  # the ragged row did not get padded
    assert evaluate_criteria({"max_count": {"errors": 0}}, out)


def test_empty_cells_become_null(ctx):
    out = validate_types({"header": ["A", "B"], "grid": [["1", ""]]}, ctx)
    assert out["typed_grid"][0][1] is None


# --- template matching -----------------------------------------------------


def typed():
    return {
        "typed_grid": [["A-100", 3], ["A-101", 10]],
        "columns": [{"name": "SKU", "type": "string"}, {"name": "Qty", "type": "integer"}],
    }


def test_template_match_normalises_column_names(ctx):
    out = map_to_schema_template(
        {**typed(), "target_schema": {"properties": {"sku": {}, "qty": {}}}}, ctx
    )
    assert out["rows"] == [{"sku": "A-100", "qty": 3}, {"sku": "A-101", "qty": 10}]


def test_template_match_accepts_a_bare_field_list(ctx):
    out = map_to_schema_template({**typed(), "target_schema": ["SKU", "Qty"]}, ctx)
    assert out["rows"][0] == {"SKU": "A-100", "Qty": 3}


def test_template_match_escalates_rather_than_emitting_holes(ctx):
    """
    Raising is how a stage says "escalate". Returning rows with a missing
    field would look like success to everything downstream.
    """
    with pytest.raises(RunnerError, match="unit_price"):
        map_to_schema_template(
            {**typed(), "target_schema": {"properties": {"sku": {}, "unit_price": {}}}}, ctx
        )


def test_optional_target_fields_do_not_force_an_escalation(ctx):
    out = map_to_schema_template(
        {
            **typed(),
            "target_schema": {
                "properties": {"sku": {}, "notes": {}},
                "required": ["sku"],
            },
        },
        ctx,
    )
    assert out["rows"][0] == {"sku": "A-100"}


# --- applying a model-produced mapping ------------------------------------


def test_apply_column_mapping_moves_values_into_place(ctx):
    out = apply_column_mapping(
        {
            **typed(),
            "target_schema": {"properties": {"part_number": {}, "quantity": {}}},
            "mapping": [
                {"target_field": "part_number", "source_column": 0},
                {"target_field": "quantity", "source_column": 1},
            ],
        },
        ctx,
    )
    assert out["rows"][0] == {"part_number": "A-100", "quantity": 3}


def test_apply_column_mapping_rejects_an_out_of_range_column(ctx):
    with pytest.raises(RunnerError, match="column 9"):
        apply_column_mapping(
            {
                **typed(),
                "target_schema": ["part_number"],
                "mapping": [{"target_field": "part_number", "source_column": 9}],
            },
            ctx,
        )


def test_apply_column_mapping_rejects_a_missing_required_field(ctx):
    with pytest.raises(RunnerError, match="omits required"):
        apply_column_mapping(
            {
                **typed(),
                "target_schema": ["part_number", "quantity"],
                "mapping": [{"target_field": "part_number", "source_column": 0}],
            },
            ctx,
        )


# --- writing ---------------------------------------------------------------


def test_write_xlsx_produces_a_readable_workbook(ctx):
    openpyxl = pytest.importorskip("openpyxl")

    out = write_xlsx({"rows": [{"sku": "A-100", "qty": 3}, {"sku": "A-101", "qty": 10}]}, ctx)
    assert evaluate_criteria({"file_exists": ["path"]}, out) == []
    assert out["row_count"] == 2

    book = openpyxl.load_workbook(out["path"])
    sheet = book.active
    assert [c.value for c in sheet[1]] == ["sku", "qty"]
    assert [c.value for c in sheet[2]] == ["A-100", 3]


def test_write_xlsx_column_order_is_stable(ctx):
    pytest.importorskip("openpyxl")
    rows = [{"b": 1, "a": 2}, {"a": 3, "b": 4, "c": 5}]
    out = write_xlsx({"rows": rows}, ctx)

    import openpyxl

    sheet = openpyxl.load_workbook(out["path"]).active
    # First row's key order, then anything later rows introduce.
    assert [c.value for c in sheet[1]] == ["b", "a", "c"]


def test_write_xlsx_refuses_an_empty_result(ctx):
    with pytest.raises(RunnerError, match="no rows"):
        write_xlsx({"rows": []}, ctx)
